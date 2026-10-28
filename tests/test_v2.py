import base64,json,threading,time,tempfile,unittest,copy,smtplib
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from v2_model import *
from v2_engine import Controller,Results
from test_independent import FakeSMTP

class FastLease:
    def __init__(self,sender):self.next_after=0;self.interval=None
    def __enter__(self):return self
    def __exit__(self,*args):pass
    def wait(self,stop):
        if stop.is_set():raise InterruptedError('simulation stopped')
    def reserve(self,interval=60):self.interval=interval;self.next_after=time.time()+.4
    def finished(self,interval=60):self.interval=interval;self.next_after=time.time()+.4

class V2Tests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);self.home=self.root/'home';self.source=self.root/'customers.xlsx'
        b=Workbook();s=b.active;s.title='Clients';s.append(['公司名称','邮箱','备注','语言','电话']);s.append(['Example One','one@example.com','Keep me','en','0100']);s.append(['Example Two','two@example.com','','es','0200']);b.save(self.source);b.close()
        self.cfg=dict(DEFAULT,market='Example Market',sender='sender@example.org',sender_name='Example',host='smtp.example.org',port=465,sheet='Clients',language_col='',storage=str(self.home))
        self.ids={}
        for lang in ('en','es','pt'):self.ids[lang]=save_template(dict(name='Example '+lang,language=lang,subject='Hello {{company}}',format='text',body='Dear {{company}},\n\nLocal example.',attachments=[]),self.home)
        self.cfg['template_id']=self.ids['en'];self.cfg=import_customers(self.source,self.cfg,self.home);self.smtp=[]
        def connector(cfg,password):
            s=FakeSMTP();self.smtp.append(s);return s
        self.controller=Controller(self.home,connector=connector,lease_factory=FastLease,copy_checker=lambda *args:'模拟副本检查，未访问网络')
        self.addCleanup(self.cleanup)
    def cleanup(self):
        if self.controller.busy:self.controller.stop();self.controller.thread.join(3)
        self.controller.close()
    def wait(self,predicate,timeout=4):
        end=time.time()+timeout
        while time.time()<end:
            if predicate():return
            time.sleep(.02)
        self.fail('Timed out waiting for simulation state')
    def finish(self):self.controller.thread.join(5);self.assertFalse(self.controller.busy)
    def test_actual_baseline_and_secret_whitelist(self):
        self.assertEqual((DEFAULT['interval'],DEFAULT['batch'],DEFAULT['cap']),(60,10,100));save_config(dict(self.cfg,password='DO_NOT_SAVE',access_token='DO_NOT_SAVE'),self.home)
        content=(self.home/'config/settings.json').read_text('utf-8');self.assertNotIn('DO_NOT_SAVE',content);self.assertNotIn('password',json.loads(content))
        with self.assertRaisesRegex(ValueError,'整数'):validate_account(dict(self.cfg,port='not-a-number'))
    def test_import_creates_stable_copy_before_sending(self):
        before=self.source.read_bytes();rows=self.controller.prepare(self.cfg);p=Path(self.controller.job['output']);self.assertTrue(p.exists());self.assertEqual(len(rows),2)
        again=import_customers(self.source,self.cfg,self.home);self.assertEqual(again['file'],self.cfg['file']);self.assertEqual(self.source.read_bytes(),before)
    def test_language_matches_existing_local_templates(self):
        cfg=dict(self.cfg,language_col='语言');rows=self.controller.prepare(cfg);self.assertEqual([r['template_id'] for r in rows],[self.ids['en'],self.ids['es']])
    def test_unknown_language_held(self):
        b=load_workbook(self.cfg['file']);b['Clients']['D3']='not-configured';b.save(self.cfg['file']);b.close();rows=self.controller.prepare(dict(self.cfg,language_col='语言'));self.assertEqual(rows[1]['state'],'template_error')
    def test_unknown_and_missing_variables_block(self):
        item=templates(self.home)[self.ids['en']];c=Customer(2,'Example','one@example.com')
        for text in ('{{missing}}','{{note}}','{{company | upper}}','{% include x %}'):
            with self.assertRaises(ValueError):build_message(c,{},self.cfg,dict(item,body=text))
    def test_all_customer_columns_and_html_escape(self):
        self.assertEqual(merge('{{电话}}',{'电话':'0100'}),'0100');self.assertEqual(merge('{{company}}',{'company':'A&B'},True),'A&amp;B')
    def test_attachment_and_inline_image_share_correct_mime(self):
        folder=self.home/'sig';folder.mkdir();png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jD4kAAAAASUVORK5CYII=');(folder/'logo.png').write_bytes(png);(folder/'sig.html').write_text('<div>Example signature<img src="logo.png" width="100"></div>','utf-8')
        signature=import_html(folder/'sig.html',self.home);item=templates(self.home)[self.ids['en']];msg=build_message(Customer(2,'Example','one@example.com'),{},dict(self.cfg,signature_html=signature),dict(item,attachments=[str(folder/'logo.png')]))
        from v2_delivery import validate_inline
        self.assertTrue(validate_inline(msg)['valid']);self.assertEqual(msg.get_content_type(),'multipart/mixed');self.assertEqual(len(list(msg.iter_attachments())),1);self.assertEqual(msg.get_body(preferencelist=('plain',)).get_content().count('Example signature'),1)
    def test_changed_attachment_blocks_start(self):
        p=self.root/'attachment.txt';p.write_text('before','utf-8');items=templates(self.home);items[self.ids['en']]['attachments']=[str(p)];atomic_json(self.home/'config/templates.json',items);self.controller.prepare(self.cfg);p.write_text('after','utf-8')
        with self.assertRaises(ValueError):self.controller.start('simulated')
    def test_pause_resume_does_not_interrupt_current_submission(self):
        self.controller.prepare(self.cfg);self.controller.start('simulated');self.wait(lambda:self.controller.snapshot()['counts']['accepted']==1);self.controller.pause();self.wait(lambda:self.controller.status=='paused');time.sleep(.5);self.assertEqual(sum(s.calls for s in self.smtp),1)
        self.controller.resume();self.finish();self.assertEqual(sum(s.calls for s in self.smtp),2);self.assertEqual(self.controller.snapshot()['counts']['accepted'],2)
    def test_stop_then_restart_never_auto_sends(self):
        self.controller.prepare(self.cfg);self.controller.start('simulated');self.wait(lambda:self.controller.snapshot()['counts']['accepted']==1);self.controller.stop();self.finish();self.controller.close()
        self.controller=Controller(self.home,connector=lambda *args:FakeSMTP(),lease_factory=FastLease,copy_checker=lambda *args:'模拟');self.assertFalse(self.controller.busy);self.assertEqual(self.controller.snapshot()['counts']['accepted'],1)
        self.controller.start('simulated');self.finish();self.assertEqual(self.controller.snapshot()['counts']['accepted'],2)
    def test_double_start_and_multi_window_blocked(self):
        self.controller.prepare(self.cfg);self.controller.start('simulated')
        with self.assertRaises(ValueError):self.controller.start('simulated')
        with self.assertRaises(ValueError):Controller(self.home)
        self.finish();self.assertEqual(sum(s.calls for s in self.smtp),2)
    def test_cross_batch_and_reimport_no_duplicates(self):
        cfg=dict(self.cfg,batch=1);self.controller.prepare(cfg);self.controller.start('simulated');self.finish();self.assertEqual(self.controller.status,'batch_done');self.controller.start('simulated');self.finish()
        rows=self.controller.prepare(cfg);self.assertTrue(all(r['state']=='history' for r in rows));self.assertEqual(sum(s.calls for s in self.smtp),2)
    def test_daily_cap_does_not_reset(self):
        self.controller.prepare(dict(self.cfg,cap=1));self.controller.start('simulated');self.finish();self.assertEqual(sum(s.calls for s in self.smtp),1);self.assertEqual(self.controller.snapshot()['today'],1)
        self.controller.start('simulated');self.finish();self.assertEqual(sum(s.calls for s in self.smtp),1)
    def test_unknown_pauses_and_is_never_retried(self):
        self.controller.connector=lambda *args:FakeSMTP('unknown');self.controller.prepare(self.cfg);self.controller.start('simulated');self.finish();self.assertEqual(self.controller.snapshot()['counts']['unknown'],1);self.assertEqual(self.controller.snapshot()['counts']['ready'],1)
        rows=self.controller.prepare(self.cfg);self.assertEqual(rows[0]['state'],'history')
    def test_failed_pauses_without_yellow(self):
        self.controller.connector=lambda *args:FakeSMTP('failed');self.controller.prepare(self.cfg);self.controller.start('simulated');self.finish();self.assertEqual(self.controller.snapshot()['counts']['failed'],1)
        b=load_workbook(self.controller.job['output']);self.assertNotEqual(b['Clients']['A2'].fill.fgColor.rgb,'FFFFFF00');self.assertIn('发送失败',b['Clients']['C2'].value);b.close()
    def test_excel_failure_does_not_change_accepted(self):
        self.controller.prepare(self.cfg)
        with patch('v2_engine.Results.record',side_effect=PermissionError('simulated')):
            self.controller.start('simulated');self.finish()
        self.assertEqual(self.controller.snapshot()['counts']['accepted'],1);self.assertEqual(sum(s.calls for s in self.smtp),1);self.controller.repair();b=load_workbook(self.controller.job['output']);self.assertEqual(b['Clients']['A2'].fill.fgColor.rgb,'FFFFFF00');b.close()
    def test_orphaned_claim_recovered_unknown_not_ready(self):
        self.controller.prepare(self.cfg);row=self.controller.job['rows'][0];self.controller.ledger.claim(self.cfg['sender'],Customer(row['row'],row['company'],row['email']),'subject',100,'<orphan@example.org>');self.controller.close();self.controller=Controller(self.home,lease_factory=FastLease)
        self.assertEqual(self.controller.snapshot()['counts']['unknown'],1);self.assertFalse(self.controller.busy)
    def test_external_excel_changes_are_backed_up_on_repair(self):
        self.controller.prepare(self.cfg);output=Path(self.controller.job['output']);b=load_workbook(output);b['Clients']['C2']='External change';b.save(output);b.close();self.controller.repair();self.assertEqual(len(list(output.parent.glob('*外部改动备份*.xlsx'))),1)
    def test_refuses_other_program_data_paths(self):
        other=self.root/'other-program';(other/'data').mkdir(parents=True);(other/'core.py').write_text('# mock program');p=other/'data'/'private.xlsx'
        with self.assertRaises(ValueError):protect_source(p)
    def test_authentication_error_is_chinese_and_no_secret(self):
        from v2_engine import explain
        text=explain(smtplib.SMTPAuthenticationError(535,b'SECRET_DO_NOT_SHOW'))
        self.assertIn('认证失败',text);self.assertNotIn('SECRET_DO_NOT_SHOW',text)
    def test_csv_working_copy_created_and_updated(self):
        path=self.root/'sample.csv';path.write_text('公司名称,邮箱,备注\nExample,one@example.com,Keep\n',encoding='utf-8');cfg=import_customers(path,dict(self.cfg,sheet='CSV'),self.home);before=path.read_bytes();self.controller.prepare(cfg)
        self.assertTrue(Path(self.controller.job['output']).is_file());self.controller.start('simulation');self.finish();text=Path(self.controller.job['output']).read_text('utf-8-sig');self.assertIn('发送邮件开发信',text);self.assertEqual(before,path.read_bytes())
    def test_auth_failure_stops_after_one_connection(self):
        calls=[]
        def reject(*args):calls.append(1);raise smtplib.SMTPAuthenticationError(535,b'FAKE_SECRET_MUST_NOT_LOG')
        self.controller.connector=reject;self.controller.prepare(self.cfg);self.controller.start('simulation');self.finish();self.assertEqual(len(calls),1);self.assertEqual(self.controller.snapshot()['counts']['accepted'],0);self.assertNotIn('FAKE_SECRET', ''.join(self.controller.logs))
    def test_throttle_421_stops_batch_with_reason(self):
        class Throttled(FakeSMTP):
            def data(self,data):raise smtplib.SMTPDataError(421,b'throttle simulation')
        self.controller.connector=lambda *args:Throttled();self.controller.prepare(self.cfg);self.controller.start('simulation');self.finish();self.assertEqual(self.controller.snapshot()['counts']['failed'],1);self.assertEqual(self.controller.snapshot()['counts']['ready'],1);self.assertIn('421',self.controller.job['rows'][0]['reason'])
    def test_language_mapping_typo_does_not_silently_fall_back(self):
        with self.assertRaises(ValueError):self.controller.prepare(dict(self.cfg,language_col='不存在的语言表头'))
    def test_ambiguous_language_templates_not_guessed(self):
        save_template(dict(name='Another Spanish',language='es',subject='Another',format='text',body='Other content',attachments=[]),self.home)
        rows=self.controller.prepare(dict(self.cfg,language_col='语言'));self.assertEqual(rows[1]['state'],'template_error');self.assertIn('多个模板',rows[1]['reason'])
    def test_literal_field_names_with_spaces_are_supported(self):
        self.assertEqual(merge('Hello {{ Contact Name }}',{'Contact Name':'Example Person'}),'Hello Example Person')
        with self.assertRaises(ValueError):merge('{{ company.upper() }}',{'company':'Example'})
    def test_unreferenced_blank_or_duplicate_extra_headers_do_not_change_import_rules(self):
        b=load_workbook(self.cfg['file']);s=b['Clients'];s['F1']='';s['G1']='';s['H1']='电话';s['H2']='duplicate';b.save(self.cfg['file']);b.close()
        rows=self.controller.prepare(self.cfg);self.assertEqual([r['state'] for r in rows],['ready','ready'])
        with self.assertRaises(ValueError):merge('{{电话}}',customer_fields(self.cfg)[2])

if __name__=='__main__':unittest.main()
