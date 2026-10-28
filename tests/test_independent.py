import base64,hashlib,json,smtplib,tempfile,threading,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from openpyxl.styles import PatternFill
import core,desktop_runtime as runtime,first_run
from core import Customer,Ledger,read_customers,render_message
from credentials import CredentialStore
from account_guard import AccountLease,require_lease
from mime_check import require_valid_inline_mime

class FakeSMTP:
    def __init__(self,state='accepted'):self.state=state;self.calls=0
    def data(self,data):return 250,b'queued local-test-only'
    def quit(self):pass
    def close(self):pass
    def sendmail(self,sender,targets,raw):
        self.calls+=1
        if self.state=='failed':raise smtplib.SMTPRecipientsRefused({targets[0]:(550,b'rejected')})
        if self.state=='unknown':raise TimeoutError()
        self.data(raw)
        if self.state=='accepted_then_disconnect':raise ConnectionError()
        return {}

class IndependentTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        for module in (runtime,first_run):
            p=patch.object(module,'BASE',self.root);p.start();self.addCleanup(p.stop)
        self.ledger=Ledger(self.root/'history.sqlite3');self.customer=Customer(2,'Example Company','customer@example.com')
        self.cfg=dict(sender='sender@example.org',cap=100,company_col='公司名称',email_col='邮箱',note_col='备注',sheet='客户')
    def message(self):return render_message(self.customer,'Hello {{company}}','Dear {{company}},\n\nBody.','sender@example.org','Example Sender')
    def test_accepted_and_duplicate_blocked(self):
        smtp=FakeSMTP();r=runtime.submit(smtp,self.ledger,self.customer,self.message(),self.cfg)
        self.assertEqual(r['smtp_code'],250);self.assertEqual(r['state'],'accepted')
        self.assertEqual(runtime.submit(smtp,self.ledger,self.customer,self.message(),self.cfg)['state'],'skipped');self.assertEqual(smtp.calls,1)
        self.assertTrue((self.root/'data/attempts/1.eml').exists())
    def test_unknown_is_not_retried(self):
        smtp=FakeSMTP('unknown');self.assertEqual(runtime.submit(smtp,self.ledger,self.customer,self.message(),self.cfg)['state'],'unknown')
        self.assertEqual(runtime.submit(smtp,self.ledger,self.customer,self.message(),self.cfg)['state'],'skipped');self.assertEqual(smtp.calls,1)
    def test_explicit_failure_recorded(self):
        smtp=FakeSMTP('failed');self.assertEqual(runtime.submit(smtp,self.ledger,self.customer,self.message(),self.cfg)['state'],'failed')
        self.assertEqual(runtime.submit(smtp,self.ledger,self.customer,self.message(),self.cfg)['state'],'skipped');self.assertEqual(smtp.calls,1)
    def test_disconnect_after_250_stays_accepted(self):
        self.assertEqual(runtime.submit(FakeSMTP('accepted_then_disconnect'),self.ledger,self.customer,self.message(),self.cfg)['state'],'accepted')
    def test_tests_do_not_block_real_recipient(self):
        r=runtime.submit(FakeSMTP(),self.ledger,self.customer,self.message(),self.cfg,test=True)
        self.assertTrue(r['is_test']);self.assertNotIn(self.customer.email,self.ledger.blocked())
    def test_daily_limit_counts_test_and_failed(self):
        cfg=dict(self.cfg,cap=1);runtime.submit(FakeSMTP('failed'),self.ledger,self.customer,self.message(),cfg,test=True)
        with self.assertRaises(ValueError):runtime.submit(FakeSMTP(),self.ledger,self.customer,self.message(),cfg)
    def source(self):
        p=self.root/'input.xlsx';b=Workbook();s=b.active;s.title='客户';s.append(['公司名称','邮箱','备注']);s.append(['Example Company','customer@example.com','有用备注']);s.append(['Other','other@example.com','']);s.column_dimensions['A'].width=31;b.save(p);b.close();return p
    def test_excel_accepted_only_actual_date_and_preservation(self):
        p=self.source();before=p.read_bytes();report=runtime.ExcelResults(dict(self.cfg,file=str(p)))
        r=dict(row=2,company='Example Company',state='accepted',finished='2026-11-03T10:00:00+08:00',message_id='<one@example.org>',smtp_code=250,smtp_response='queued')
        report.record(r);report.record(r);report.record(dict(r,row=3,state='unknown',message_id='<two@example.org>'))
        b=load_workbook(report.path);s=b['客户'];self.assertEqual(s['A2'].fill.fgColor.rgb,'FFFFFF00');self.assertNotEqual(s['A3'].fill.fgColor.rgb,'FFFFFF00')
        self.assertEqual(s['C2'].value,'有用备注\n发送邮件开发信（11.3）');self.assertIn('待核实',s['C3'].value);self.assertEqual(s.column_dimensions['A'].width,31);b.close();self.assertEqual(p.read_bytes(),before)
    def test_changed_source_does_not_overwrite(self):
        p=self.source();report=runtime.ExcelResults(dict(self.cfg,file=str(p)));p.write_bytes(b'changed')
        with self.assertRaises(ValueError):report.write()
    def test_import_filters_yellow_invalid_duplicate_and_reply(self):
        p=self.source();b=load_workbook(p);s=b['客户'];s['A2'].fill=PatternFill('solid',fgColor='FFFFFF00');s.append(['Again','customer@example.com','']);s.append(['Bad','官网表单','']);s.append(['Reply','reply@example.com','']);b.save(p);b.close()
        rows=read_customers(p,'客户','公司名称','邮箱','备注',blocked={'reply@example.com'})
        self.assertEqual([r.status for r in rows],['yellow','ready','yellow','invalid','blocked'])
    def test_duplicate_email_in_order(self):
        p=self.source();b=load_workbook(p);b['客户'].append(['Duplicate','customer@example.com','']);b.save(p);b.close()
        self.assertEqual([r.status for r in read_customers(p,'客户','公司名称','邮箱')],['ready','ready','duplicate'])
    def test_signature_import_cid_original_image_and_move(self):
        source=self.root/'original';source.mkdir();image=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jD4kAAAAASUVORK5CYII=');(source/'logo.png').write_bytes(image)
        original='<div style="color:#123456"><img src="logo.png" width="120"><a href="https://example.com">Example</a></div>';(source/'sig.html').write_text(original,'utf-8')
        target=first_run.import_signature(source/'sig.html');msg=render_message(self.customer,'Hello','Paragraph 1\n\nParagraph 2','sender@example.org','Example',signature_html=self.root/target)
        result=require_valid_inline_mime(msg.as_bytes(),[image]);self.assertTrue(result['original_image_bytes_preserved']);self.assertEqual((source/'sig.html').read_text('utf-8'),original)
        self.assertNotIn('file:',next(p.get_content() for p in msg.walk() if p.get_content_type()=='text/html'))
    def test_missing_signature_resource_rejected(self):
        p=self.root/'bad.html';p.write_text('<img src="missing.png">','utf-8')
        with self.assertRaises(ValueError):first_run.import_signature(p)
    def test_credentials_have_isolated_namespace(self):
        target=CredentialStore.target(dict(host='smtp.example.org',port=465,security='SSL/TLS',sender='sender@example.org',username=''))
        self.assertTrue(target.startswith('OpenDeskMail/SMTP/legacy/'));self.assertNotIn('sender@',target)
    def test_lease_spacing_persists_and_pause(self):
        directory=self.root/'locks'
        with AccountLease('sender@example.org',directory) as lease:
            require_lease('sender@example.org');lease.reserve()
            with self.assertRaises(RuntimeError):lease.reserve()
        with AccountLease('sender@example.org',directory) as lease:
            stop=threading.Event();stop.set()
            with self.assertRaises(InterruptedError):lease.wait(stop)
        with self.assertRaises(RuntimeError):require_lease('sender@example.org')
    def test_no_imap_does_not_claim_archived(self):
        self.assertIn('未核实',runtime.verify_copies({},'not-a-real-secret',[]))

if __name__=='__main__':unittest.main()
