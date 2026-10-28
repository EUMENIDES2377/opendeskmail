"""Durable wire results, incremental Excel output and read-only Sent checks."""
import hashlib,json,os,smtplib
from datetime import datetime
from email import policy
from pathlib import Path
from core import BASE,now

def validate_inline(msg):
    import copy
    from mime_check import require_valid_inline_mime
    def remove_attachments(part):
        if part.is_multipart():
            part.set_payload([p for p in part.get_payload() if p.get_content_disposition()!='attachment'])
            for child in part.get_payload():remove_attachments(child)
    clean=copy.deepcopy(msg);remove_attachments(clean);return require_valid_inline_mime(clean.as_bytes(policy=policy.SMTP))

def atomic_json(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);temp=path.with_suffix('.tmp')
    with temp.open('w',encoding='utf-8') as f:json.dump(value,f,ensure_ascii=False,indent=2);f.flush();os.fsync(f.fileno())
    os.replace(temp,path)

def submit(smtp,ledger,c,msg,cfg,test=False):
    raw=msg.as_bytes(policy=policy.SMTP)
    if any(p.get_content_disposition()=='inline' for p in msg.walk()):
        from v2_delivery import validate_inline
        validate_inline(msg)
    ident=ledger.claim(cfg['sender'],c,str(msg['Subject']),cfg['cap'],str(msg['Message-ID']),test=test)
    if ident is None:return {'state':'skipped','row':c.row,'company':c.company,'email':c.email}
    r=dict(attempt_id=ident,row=c.row,company=c.company,email=c.email,subject=str(msg['Subject']),message_id=str(msg['Message-ID']),started=now(),state='unknown',is_test=test)
    dest=Path(cfg['storage'])/'data'/'attempts';dest.mkdir(parents=True,exist_ok=True)
    # Keep the ledger claim blocking if saving the wire intent itself fails.
    atomic_json(dest/f'{ident}-intent.json',r);(dest/f'{ident}.eml').write_bytes(raw)
    response=None;original=smtp.data
    def capture(data):
        nonlocal response
        response=original(data);return response
    smtp.data=capture
    try:
        refused=smtp.sendmail(cfg['sender'],[c.email],raw)
        if refused:raise smtplib.SMTPRecipientsRefused(refused)
        if not response or response[0]!=250:raise RuntimeError('未取得最终DATA 250响应')
        r['state']='accepted'
    except (smtplib.SMTPRecipientsRefused,smtplib.SMTPSenderRefused,smtplib.SMTPDataError) as exc:
        code=getattr(exc,'smtp_code',None)
        if isinstance(exc,smtplib.SMTPRecipientsRefused):code=next((v[0] for v in exc.recipients.values()),None)
        reason={'SMTPRecipientsRefused':'服务器拒绝收件人，请核对邮箱有效性或服务商策略','SMTPSenderRefused':'服务器拒绝发件账号，请核对发件权限','SMTPDataError':'服务器拒绝邮件内容或限流，请联系管理员核查'}[type(exc).__name__]
        r.update(state='failed',reason=reason+(f'（{code}）' if code else ''),smtp_code=code or '',smtp_response=type(exc).__name__)
    except Exception as exc:r.update(state='accepted' if response and response[0]==250 else 'unknown',reason=type(exc).__name__)
    finally:smtp.data=original
    if response:r.update(smtp_code=response[0],smtp_response=response[1].decode('utf-8','replace') if isinstance(response[1],bytes) else str(response[1]))
    r['finished']=now();state='test_accepted' if test and r['state']=='accepted' else r['state']
    ledger.finish(ident,state,('服务器已接受（未确认客户收到）' if r['state']=='accepted' else r.get('reason',''))+f"；{r.get('smtp_code','')} {r.get('smtp_response','')}")
    atomic_json(dest/f'{ident}-result.json',r);return r

def verify_copies(cfg,password,results):
    if not cfg.get('imap_host') or not cfg.get('imap_sent_folder'):return '未配置IMAP/Sent文件夹；副本未核实，本地记录已保留'
    import mail_archive
    from email.parser import BytesParser
    client=mail_archive.connect(cfg,password)
    try:
        status,_=client.select(mail_archive.quoted(cfg['imap_sent_folder']),readonly=True)
        if status!='OK':raise ValueError('无法只读打开配置的Sent文件夹')
        status,data=client.uid('fetch','1:*','(UID FLAGS BODY.PEEK[HEADER.FIELDS (MESSAGE-ID)])')
        if status!='OK':raise ValueError('副本邮件头读取失败')
        ids=[]
        for item in data:
            if isinstance(item,tuple) and b'\\Deleted' not in item[0]:ids.append(str(BytesParser(policy=policy.default).parsebytes(item[1]).get('Message-ID','')).strip())
        for r in results:
            if r['state']=='accepted':r['sent_folder_matches']=ids.count(r['message_id'])
        return '只读核查已完成；副本可能延迟出现，未核实不触发APPEND或重发'
    finally:client.logout()

class ExcelResults:
    def __init__(self,cfg):
        self.cfg=cfg;self.source=Path(cfg['file']);self.source_hash=hashlib.sha256(self.source.read_bytes()).hexdigest()
        key=hashlib.sha256((str(self.source.resolve())+'|'+self.source_hash+'|'+cfg['sheet']+'|'+cfg['sender']).encode()).hexdigest()[:20]
        self.path=Path(cfg['storage'])/'data'/'results'/(self.source.stem+'_'+key+'_发送更新.xlsx');self.state=Path(cfg['storage'])/'data'/'results'/(key+'.json')
        self.rows=json.loads(self.state.read_text('utf-8')).get('rows',{}) if self.state.exists() else {}
    def record(self,r):
        if r.get('is_test') or r['state']=='skipped':return
        self.rows[str(r['row'])]=r;atomic_json(self.state,dict(rows=self.rows,source_sha256=self.source_hash,output=str(self.path)))
        self.write()
    def write(self):
        if self.source.suffix.lower()!='.xlsx':return
        if hashlib.sha256(self.source.read_bytes()).hexdigest()!=self.source_hash:raise ValueError('原客户表已变化，已保存发送日志并暂停，禁止覆盖或重发')
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill
        book=load_workbook(self.source)
        try:
            sh=book[self.cfg['sheet']];heads=[str(c.value or '').strip() for c in sh[1]];ci=heads.index(self.cfg['company_col'])+1
            def column(label):
                if label not in heads:heads.append(label);sh.cell(1,len(heads),label);sh.column_dimensions[sh.cell(1,len(heads)).column_letter].width=38
                return heads.index(label)+1
            ni=column(self.cfg['note_col'] or '备注');si=column('本工具发送状态');ti=column('本工具发送时间');mi=column('本工具Message-ID');ri=column('本工具SMTP回复');ai=column('本工具副本核对')
            for r in self.rows.values():
                n=r['row'];dt=datetime.fromisoformat(r['finished']);day=f'{dt.month}.{dt.day}';state=r['state'];note=''
                if state=='accepted':sh.cell(n,ci).fill=PatternFill('solid',fgColor='FFFFFF00');note=f'发送邮件开发信（{day}）'
                elif state=='failed':note=f'发送失败（{day}）：'+r.get('reason','')
                elif state=='unknown':note=f'发送状态待核实（{day}）'
                existing=str(sh.cell(n,ni).value or '')
                if note and note not in existing.splitlines():sh.cell(n,ni,note if not existing.strip() or existing.strip()=='待发送邮件开发信' else existing.rstrip()+'\n'+note)
                sh.cell(n,si,{'accepted':'SMTP服务器已接受，未确认客户收到','failed':'发送失败','unknown':'发送状态待核实'}[state]);sh.cell(n,ti,dt.replace(tzinfo=None)).number_format='yyyy-mm-dd hh:mm:ss'
                sh.cell(n,mi,r.get('message_id',''));sh.cell(n,ri,str(r.get('smtp_code',''))+' '+r.get('smtp_response',''));sh.cell(n,ai,'已核实'+str(r['sent_folder_matches'])+'份' if r.get('sent_folder_matches') else '未核实（不重发）')
            self.path.parent.mkdir(parents=True,exist_ok=True);temp=self.path.with_suffix('.tmp.xlsx');book.save(temp);os.replace(temp,self.path)
        finally:book.close()
