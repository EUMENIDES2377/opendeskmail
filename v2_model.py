"""Local-only configuration, imported workbooks and deterministic MIME templates."""
import copy,hashlib,html,json,mimetypes,re,shutil,tempfile
from pathlib import Path
from email.message import EmailMessage
from email.utils import formataddr,formatdate,make_msgid
from core import BASE,Customer,read_customers,normalize,valid_email
from html_signature import load_signature,body_html,plain_text,decode_html
from desktop_runtime import atomic_json
from credentials import CredentialStore

from app_paths import USER_HOME
HOME=USER_HOME
DEFAULT=dict(market='',sender='',sender_name='',host='',port='',security='SSL/TLS',username='',
             interval=60,batch=10,cap=100,imap_host='',imap_port=993,imap_security='SSL/TLS',imap_sent_folder='',
             signature_html='',signature_text='',template_id='',sheet='',company_col='公司名称',email_col='邮箱',note_col='备注',language_col='',credential_kind='password')
STATES={'ready':'待发送','accepted':'服务器已接受','failed':'发送失败','unknown':'状态待核实','sending':'正在提交','skipped':'跳过',
        'history':'已有记录','yellow':'已标黄','duplicate':'重复邮箱','invalid':'无效字段','blocked':'不再联系/已回复','template_error':'模板或变量缺失'}

def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def public_config(value):return {k:v for k,v in value.items() if k in DEFAULT or k in ('file','storage','import_id')}
def load_config(root=HOME):
    root=Path(root);path=root/'config/settings.json';cfg=DEFAULT.copy()
    if path.exists():cfg.update(public_config(json.loads(path.read_text('utf-8'))))
    return cfg
def validate_account(cfg):
    if not cfg.get('market') or not valid_email(normalize(cfg.get('sender'))) or not cfg.get('host'):raise ValueError('请填写项目、有效邮箱和SMTP服务器。')
    if cfg.get('security') not in ('SSL/TLS','STARTTLS'):raise ValueError('必须选择SSL/TLS或STARTTLS加密。')
    if cfg.get('credential_kind','password') not in ('password','smtp_authorization_code'):raise ValueError('凭据类型须为邮箱登录密码或SMTP授权码；OAuth需另行适配。')
    try:
        for key in ('port','imap_port','interval','batch','cap'):cfg[key]=int(cfg.get(key) or (993 if key=='imap_port' else 0))
    except (ValueError,TypeError):raise ValueError('端口、间隔、批次数和每日上限必须填写整数，请检查输入。') from None
    if not 1<=cfg['port']<=65535 or cfg['interval']<60 or cfg['batch']<1 or cfg['cap']<0:raise ValueError('端口1–65535；间隔至少60秒；每批至少1封；每日上限0表示不限。')
    if cfg.get('imap_host') and (cfg.get('imap_security')!='SSL/TLS' or not 1<=cfg['imap_port']<=65535):raise ValueError('IMAP仅支持经过证书验证的SSL/TLS连接。')
    cfg['sender']=normalize(cfg['sender']);return cfg
def save_config(cfg,root=HOME):atomic_json(Path(root)/'config/settings.json',public_config(cfg))

class SecureStore(CredentialStore):
    @staticmethod
    def target(cfg):return 'OpenDeskMail/SMTP/v1/'+CredentialStore.target(cfg).rsplit('/',1)[-1]

def protect_source(path,app_root=BASE):
    path=Path(path).resolve()
    for ancestor in path.parents:
        if ancestor==Path(app_root).resolve():continue
        if (ancestor/'core.py').is_file() and any(p.lower() in {'data','logs','verification','workspace','.venv'} for p in path.relative_to(ancestor).parts):
            raise ValueError('所选路径属于其他邮件程序的数据目录。请先另存普通客户表副本，再导入；本软件不会写入该目录。')
    return path

def import_customers(path,cfg,root=HOME):
    root=Path(root);source=protect_source(path)
    if source.suffix.lower() not in ('.xlsx','.csv'):raise ValueError('仅支持xlsx或UTF-8 CSV。')
    key=hashlib.sha256((str(source)+'|'+digest(source)).encode()).hexdigest()[:24];folder=root/'imports'/key;folder.mkdir(parents=True,exist_ok=True)
    snapshot=folder/('source'+source.suffix.lower())
    if not snapshot.exists():shutil.copyfile(source,snapshot)
    if digest(snapshot)!=digest(source):raise ValueError('导入快照校验失败，请检查副本。')
    metadata=dict(original=str(source),source_sha256=digest(snapshot),snapshot=str(snapshot))
    atomic_json(folder/'import.json',metadata)
    return dict(cfg,file=str(snapshot),storage=str(root),import_id=key)

LANG={'english':'en','英语':'en','en':'en','español':'es','spanish':'es','西语':'es','西班牙语':'es','es':'es','portuguese':'pt','português':'pt','葡语':'pt','葡萄牙语':'pt','pt':'pt'}
def language(value):return LANG.get(str(value).strip().casefold(),str(value).strip().casefold())
def templates(root=HOME):
    p=Path(root)/'config/templates.json';return json.loads(p.read_text('utf-8')) if p.exists() else {}
def save_template(value,root=HOME):
    root=Path(root);all_items=templates(root);ident=value.get('id') or hashlib.sha256((value['name']+'|'+language(value['language'])).encode()).hexdigest()[:16]
    value=dict(value,id=ident,language=language(value['language']));all_items[ident]=value;atomic_json(root/'config/templates.json',all_items);return ident

def import_html(path,root=HOME,kind='signature'):
    sig=load_signature(path);folder=Path(root)/'assets'/kind/sig.fingerprint[:20];folder.mkdir(parents=True,exist_ok=True);markup=sig.markup
    for asset in sig.images:
        name=asset.cid.split('@')[0]+(mimetypes.guess_extension(asset.media_type) or '.bin');(folder/name).write_bytes(asset.data);markup=markup.replace('cid:'+asset.cid,name)
    target=folder/'content.html';target.write_text('<meta charset="utf-8">'+markup,'utf-8');return str(target)
def paste_html(markup,resource_dir,root=HOME,kind='signature'):
    # Resolve pasted relative resources against only the directory selected by user.
    folder=Path(root)/'drafts';folder.mkdir(parents=True,exist_ok=True)
    from html_signature import map_resources,resolve_local
    def absolute(url):
        if url.startswith('data:image/'):return url
        return resolve_local(url,Path(resource_dir or folder)).as_uri()
    markup=map_resources(markup,absolute);p=folder/(kind+'.html');p.write_text(markup,'utf-8');return import_html(p,root,kind)

VARIABLE=re.compile(r'{{\s*([^{}\r\n]+?)\s*}}',re.UNICODE)
def merge(text,fields,escape=False):
    def replace(match):
        key=match.group(1).strip()
        if key not in fields or fields[key] is None or str(fields[key]).strip()=='':raise ValueError('缺少模板字段：'+key)
        value=str(fields[key]);return html.escape(value,quote=True) if escape else value
    result=VARIABLE.sub(replace,text)
    if any(t in result for t in ('{{','}}','{%','{#')):raise ValueError('存在未知变量或不支持的模板表达式；只允许 {{字段名}}。')
    return result

def customer_fields(cfg):
    source=Path(cfg['file']);records={}
    if source.suffix.lower()=='.csv':
        import csv
        with source.open(encoding='utf-8-sig',newline='') as f:
            reader=csv.DictReader(f)
            if cfg.get('language_col') and cfg['language_col'] not in (reader.fieldnames or []):raise ValueError('所配置的语言列表头不存在。')
            for row,item in enumerate(reader,2):records[row]=item
    else:
        from openpyxl import load_workbook
        b=load_workbook(source,data_only=False)
        try:
            sh=b[cfg['sheet']];headers=[str(c.value or '').strip() for c in sh[1]]
            duplicates={name for name in headers if name and headers.count(name)>1}
            if cfg.get('language_col') and cfg['language_col'] not in headers:raise ValueError('所配置的语言列表头不存在。')
            for row in sh.iter_rows(min_row=2):records[row[0].row]={headers[i]:c.value if c.data_type!='f' and headers[i] not in duplicates else None for i,c in enumerate(row) if headers[i]}
        finally:b.close()
    return records

def build_message(c,fields,cfg,template):
    fields=dict(fields,company=c.company,email=c.email,note=c.note,sender=cfg['sender'])
    subject=merge(template['subject'],fields)
    if not subject.strip() or '\r' in subject or '\n' in subject:raise ValueError('主题不能为空或含换行。')
    if not valid_email(cfg['sender']) or not valid_email(c.email):raise ValueError('发件人或收件人邮箱无效。')
    images={};body=template.get('body','');markup=None
    if template.get('format')=='html':
        loaded=load_signature(template['html_path']);markup=merge(loaded.markup,fields,True);body=plain_text(markup);images.update({x.cid:x for x in loaded.images})
    else:body=merge(body,fields)
    if not body.strip():raise ValueError('正文不能为空。')
    signature=merge(cfg.get('signature_text',''),fields)
    if cfg.get('signature_html'):
        sig=load_signature(cfg['signature_html']);signature=merge(sig.text,fields)
        markup=(markup if markup is not None else body_html(body))+merge(sig.markup,fields,True);images.update({x.cid:x for x in sig.images})
    elif signature:
        if markup is not None:markup+=body_html(signature)
    msg=EmailMessage();msg['From']=formataddr((cfg.get('sender_name',''),cfg['sender']));msg['To']=c.email;msg['Subject']=subject;msg['Date']=formatdate(localtime=False);msg['Message-ID']=make_msgid(domain=cfg['sender'].split('@')[1])
    msg.set_content(body.rstrip()+ ('\n\n'+signature if signature else '')+'\n')
    if markup is not None:
        msg.add_alternative(markup,subtype='html');part=msg.get_payload()[-1]
        for asset in images.values():
            major,minor=asset.media_type.split('/',1);part.add_related(asset.data,maintype=major,subtype=minor,cid='<'+asset.cid+'>',disposition='inline',filename=asset.filename)
        if images:
            cid='body-'+hashlib.sha256(str(msg['Message-ID']).encode()).hexdigest()[:24]+'@independent-mail.local';part.get_payload()[0]['Content-ID']='<'+cid+'>';part.set_param('type','text/html');part.set_param('start','<'+cid+'>')
    for path in template.get('attachments',[]):
        path=Path(path)
        if not path.is_file():raise ValueError('附件不存在：'+path.name)
        media=mimetypes.guess_type(path.name)[0] or 'application/octet-stream';major,minor=media.split('/',1);msg.add_attachment(path.read_bytes(),maintype=major,subtype=minor,filename=path.name)
    return msg

def template_fingerprint(all_items,cfg):
    hashes=[]
    for item in all_items.values():
        hashes.append(item)
        if item.get('format')=='html':hashes.append(load_signature(item['html_path']).fingerprint)
        hashes.extend(digest(p) for p in item.get('attachments',[]))
    if cfg.get('signature_html'):hashes.append(load_signature(cfg['signature_html']).fingerprint)
    return hashlib.sha256(json.dumps(hashes,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
