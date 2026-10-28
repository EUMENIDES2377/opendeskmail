"""Excel adapter and durable sending queue around awdeorio/mailmerge 2.2.4."""
from __future__ import annotations
import csv
import json
import re
import smtplib
import sqlite3
import ssl
import tempfile
import hashlib
from dataclasses import dataclass, asdict
from contextlib import contextmanager
from datetime import datetime
from email.utils import formataddr, make_msgid, formatdate
from email import policy
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

BASE = Path(__file__).resolve().parent
ADDRESS = re.compile(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?\.[A-Za-z]{2,}")
STATUSES = {"ready":"待发送", "yellow":"已标黄，跳过", "history":"已有记录，跳过",
            "blocked":"不再联系，跳过", "duplicate":"表内重复，跳过", "invalid":"邮箱或公司名无效"}

def now():
    return datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')

def normalize(value):
    return str(value or '').strip().lower()

def valid_email(value):
    return bool(ADDRESS.fullmatch(value))

def one_line(value):
    return re.sub(r'[\r\n\x00-\x1f]+', ' ', str(value or '')).strip()

def address_file(path):
    if not Path(path).exists():
        return set()
    return {normalize(line) for line in Path(path).read_text(encoding='utf-8-sig').splitlines()
            if line.strip() and not line.lstrip().startswith('#')}

def is_yellow(cell):
    """Explicit solid RGB/indexed fills only; conditional formatting is not evaluated."""
    if cell.fill.patternType != 'solid':
        return False
    c = cell.fill.fgColor
    if c.type == 'indexed':
        from openpyxl.styles.colors import COLOR_INDEX
        rgb = COLOR_INDEX[c.indexed] if c.indexed < len(COLOR_INDEX) else ''
    elif c.type == 'rgb':
        rgb = c.rgb
    else:
        return False
    try:
        r, g, b = (int(rgb[-6:][i:i+2], 16) for i in (0, 2, 4))
        return r >= 210 and g >= 180 and b <= 160
    except (TypeError, ValueError):
        return False

@dataclass
class Customer:
    row: int
    company: str
    email: str
    note: str = ''
    status: str = 'ready'

def sheet_names(path):
    if Path(path).suffix.lower() == '.csv':
        return ['CSV']
    if Path(path).suffix.lower() != '.xlsx':
        raise ValueError('支持 .xlsx 和 UTF-8 CSV；旧版 .xls 请先另存为 .xlsx。')
    from openpyxl import load_workbook
    w = load_workbook(path, read_only=True)
    try:
        return w.sheetnames
    finally:
        w.close()

def read_customers(path, sheet, company_col='公司名', email_col='邮箱', note_col='备注', history=(), blocked=()):
    source = []
    if Path(path).suffix.lower() == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            if company_col not in headers or email_col not in headers:
                raise ValueError('找不到公司名或邮箱表头；请填写与第1行完全一致的列名。')
            for n, row in enumerate(reader, 2):
                source.append((n, row.get(company_col), row.get(email_col), row.get(note_col), False))
    else:
        from openpyxl import load_workbook
        w = load_workbook(path, data_only=False)
        try:
            s = w[sheet]
            headers = [str(c.value or '').strip() for c in s[1]]
            if headers.count(company_col) != 1 or headers.count(email_col) != 1:
                raise ValueError('公司名、邮箱表头必须各出现一次，并位于第1行。')
            ci, ei = headers.index(company_col)+1, headers.index(email_col)+1
            ni = headers.index(note_col)+1 if note_col in headers else None
            for n in range(2, s.max_row+1):
                c, e = s.cell(n, ci), s.cell(n, ei)
                # Formulas in recipient/name fields are not executable customer data.
                company = '' if c.data_type == 'f' else c.value
                email = '' if e.data_type == 'f' else e.value
                source.append((n, company, email, s.cell(n, ni).value if ni else '', is_yellow(c)))
        finally:
            w.close()
    result, seen = [], set()
    history, blocked = set(history), set(blocked)
    # A yellow duplicate anywhere in this sheet suppresses unmarked copies too.
    yellow_emails = {normalize(e) for _, _, e, _, yellow in source if yellow}
    for n, c, e, note, yellow in source:
        company, email = one_line(c), normalize(e)
        if not company and not email:
            continue
        status = 'ready'
        if not company or not valid_email(email):
            status = 'invalid'
        elif email in blocked:
            status = 'blocked'
        elif yellow or email in yellow_emails:
            status = 'yellow'
        elif email in history:
            status = 'history'
        elif email in seen:
            status = 'duplicate'
        seen.add(email)
        result.append(Customer(n, company, email, one_line(note), status))
    return result

class Ledger:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS attempts (id INTEGER PRIMARY KEY, sender TEXT, email TEXT, company TEXT, subject TEXT, state TEXT, started TEXT, finished TEXT, detail TEXT, message_id TEXT)')
            columns = {row[1] for row in db.execute('PRAGMA table_info(attempts)')}
            if 'is_test' not in columns:
                db.execute('ALTER TABLE attempts ADD COLUMN is_test INTEGER NOT NULL DEFAULT 0')
            if 'test_key' not in columns:
                db.execute('ALTER TABLE attempts ADD COLUMN test_key TEXT')
            db.execute('CREATE UNIQUE INDEX IF NOT EXISTS unique_test_key ON attempts(test_key) WHERE test_key IS NOT NULL')

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=10)
        try:
            with db:
                yield db
        finally:
            db.close()

    def blocked(self):
        with self.connect() as db:
            return {r[0] for r in db.execute("SELECT email FROM attempts WHERE is_test=0 AND state IN ('sending','accepted','unknown','manual','failed')")}

    def claim(self, sender, customer, subject, cap, message_id, test=False, test_key=None):
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            if test_key and db.execute('SELECT 1 FROM attempts WHERE test_key=?', (test_key,)).fetchone():
                return None
            if not test and db.execute("SELECT 1 FROM attempts WHERE email=? AND is_test=0 AND state IN ('sending','accepted','unknown','manual','failed')", (customer.email,)).fetchone():
                return None
            count = db.execute("SELECT COUNT(*) FROM attempts WHERE sender=? AND substr(started,1,10)=? AND state!='manual'", (sender, now()[:10])).fetchone()[0]
            if cap and count >= cap:
                raise ValueError('已达到今天的发送尝试上限（北京时间，含失败和测试邮件）。')
            cur = db.execute('INSERT INTO attempts(sender,email,company,subject,state,started,message_id,is_test,test_key) VALUES(?,?,?,?,?,?,?,?,?)',
                             (sender, customer.email, customer.company, subject, 'sending', now(), message_id, int(test), test_key))
            return cur.lastrowid

    def finish(self, ident, state, detail=''):
        with self.connect() as db:
            db.execute('UPDATE attempts SET state=?,finished=?,detail=? WHERE id=?', (state, now(), detail, ident))

    def rows(self):
        with self.connect() as db:
            return list(db.execute('SELECT started,email,company,subject,state,detail,message_id FROM attempts ORDER BY id'))

    def archive_result(self, ident, detail):
        with self.connect() as db:
            db.execute("UPDATE attempts SET detail=coalesce(detail,'') || '；' || ? WHERE id=?", (detail, ident))

def render_message(customer, subject, body, sender, sender_name, signature='', signature_html=None):
    from mailmerge import TemplateMessage
    if not valid_email(sender):
        raise ValueError('请填写正确的发件邮箱。')
    if '\n' in subject or '\r' in subject or not subject.strip():
        raise ValueError('主题不能为空或包含换行。')
    if not body.strip():
        raise ValueError('正文不能为空。')
    raw = ('To: {{email}}\nFrom: {{sender}}\nSubject: ' + subject +
           '\nContent-Type: text/plain\n\n' + body + ('\n\n'+signature if signature.strip() and not signature_html else ''))
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp)/'template.txt'
        path.write_text(raw, encoding='utf-8')
        _, recipients, msg = TemplateMessage(path).render({
            'company':customer.company, 'email':customer.email, 'note':customer.note,
            'sender':formataddr((one_line(sender_name), sender)),
        })
    if recipients != [customer.email] or msg.get_all('Cc') or msg.get_all('Bcc'):
        raise ValueError('每封邮件只能包含当前客户这一个收件人。')
    msg['Message-ID'] = make_msgid(domain=sender.split('@')[1])
    if signature_html:
        from html_signature import load_signature, message_plain, body_html
        sig = load_signature(signature_html)
        rendered_body = message_plain(msg)
        multipart = EmailMessage()
        for key in ('To', 'From', 'Subject', 'Message-ID'):
            multipart[key] = msg[key]
        multipart['Date'] = msg.get('Date') or formatdate(localtime=False)
        multipart.set_content(rendered_body.rstrip() + '\n\n' + sig.text + '\n')
        multipart.add_alternative(body_html(rendered_body.rstrip()) + sig.markup, subtype='html')
        html_part = multipart.get_payload()[-1]
        for image in sig.images:
            major, minor = image.media_type.split('/', 1)
            html_part.add_related(image.data, maintype=major, subtype=minor,
                                  cid='<' + image.cid + '>', disposition='inline', filename=image.filename)
        if sig.images:
            # Explicit related root helps clients select HTML and its inline
            # resources without treating the image as an unrelated attachment.
            root_id = 'body-' + hashlib.sha256(str(msg['Message-ID']).encode()).hexdigest()[:24] + '@independent-mail.local'
            html_part.get_payload()[0]['Content-ID'] = '<' + root_id + '>'
            html_part.set_param('type', 'text/html')
            html_part.set_param('start', '<' + root_id + '>')
        return multipart
    return msg

def smtp_connect(config, password):
    from account_guard import require_lease
    require_lease(config['sender'])
    host, port, security = config['host'].strip(), int(config['port']), config['security']
    sender = normalize(config['sender'])
    if not host or not valid_email(sender) or not password:
        raise ValueError('请填写SMTP服务器、发件邮箱和密码/授权码。')
    ctx = ssl.create_default_context()
    if security == 'SSL/TLS':
        smtp = smtplib.SMTP_SSL(host, port, timeout=30, context=ctx)
    elif security == 'STARTTLS':
        smtp = smtplib.SMTP(host, port, timeout=30)
    else:
        raise ValueError('请选择SSL/TLS或STARTTLS。')
    try:
        smtp.ehlo()
        if security == 'STARTTLS':
            smtp.starttls(context=ctx)
            smtp.ehlo()
        # Restrict smtplib.login to one mechanism, so an authentication rejection
        # never causes fallback attempts with the same credential.
        features = getattr(smtp, 'esmtp_features', None)
        if isinstance(features, dict):
            methods = features.get('auth', '').upper().split()
            selected = next((m for m in ('PLAIN', 'LOGIN', 'CRAM-MD5') if m in methods), None)
            if not selected:
                raise ValueError('服务器未提供受支持的密码认证；若要求 OAuth/XOAUTH2，需要适配 OAuth，不能反复尝试密码。')
            smtp.esmtp_features['auth'] = selected
        smtp.login(config.get('username', '').strip() or sender, password)
        return smtp
    except Exception:
        smtp.close()
        raise

def send_one(smtp, ledger, customer, msg, sender, cap, test=False, test_key=None, archive=None):
    ident = ledger.claim(sender, customer, str(msg['Subject']), cap, msg['Message-ID'], test=test, test_key=test_key)
    if ident is None:
        return 'skipped'
    try:
        # Serialize before SMTP to distinguish local failures from uncertain delivery.
        data = msg.as_bytes(policy=policy.SMTP)
        if any(part.get_content_maintype() == 'image' for part in msg.walk()):
            from mime_check import require_valid_inline_mime
            require_valid_inline_mime(data)
    except Exception:
        ledger.finish(ident, 'failed', '本地邮件编码失败')
        raise
    try:
        refused = smtp.sendmail(sender, [customer.email], data)
        if refused:
            raise smtplib.SMTPRecipientsRefused(refused)
    except (smtplib.SMTPRecipientsRefused, smtplib.SMTPSenderRefused, smtplib.SMTPDataError) as exc:
        ledger.finish(ident, 'failed', f'服务器明确拒绝：{type(exc).__name__}')
        raise
    except Exception as exc:
        # No blind retry: timeout/disconnection may happen after server acceptance.
        ledger.finish(ident, 'unknown', f'结果待核查：{type(exc).__name__}')
        raise
    ledger.finish(ident, 'test_accepted' if test else 'accepted', '发件服务器已接受；不代表已送达或进入收件箱')
    if archive:
        try:
            archive(msg)
        except Exception as exc:
            ledger.archive_result(ident, '发送已接受，副本保存失败：' + type(exc).__name__ + '；禁止因此重新发送')
        else:
            ledger.archive_result(ident, 'IMAP 副本保存成功')
    else:
        ledger.archive_result(ident, '服务器已发送副本尚未核实；本工具保留本地记录')
    return 'accepted'

def safe_csv(value):
    text = str(value or '')
    return "'"+text if text.lstrip().startswith(('=', '+', '-', '@')) else text

def export_csv(path, headers, rows):
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows([[safe_csv(v) for v in row] for row in rows])
